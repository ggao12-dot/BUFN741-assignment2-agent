"""Create frozen editable DOCX/PDF deliverables from a validated run Markdown."""
from __future__ import annotations

import argparse
import csv
from datetime import datetime, timezone
import hashlib
import json
import os
from pathlib import Path
import re
import shutil
import socket
import subprocess
import sys
import time
import zipfile

from docx import Document
from docx.enum.section import WD_ORIENT, WD_SECTION
from docx.enum.style import WD_STYLE_TYPE
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


REQUIRED_SECTIONS = [
    'Executive Summary', '1. Objective and Methodology', '2. Data and Sample Construction',
    '3. Descriptive Statistics and Diversification', '4. Markowitz Full-Sample Results',
    '5. Benchmark Portfolios', '6. Effect of Portfolio Constraints',
    '7. Pre-COVID versus Post-COVID Markowitz Results', '8. Single-Index Model Estimation',
    '9. SIM Portfolio Optimization', '10. Markowitz versus SIM: Covariance-Only Experiment',
    '11. Complete Markowitz versus SIM Comparison', '12. Covariance-Model Diagnostics',
    '13. Answers to Assignment Analysis Questions', '14. Practical Interpretation and Limitations',
    '15. Conclusion', 'Appendix A. Provenance and Reproducibility']
LANDSCAPE_TABLES = {2, 5, 8, 9, 11}


def sha256(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open('rb') as handle:
        for chunk in iter(lambda: handle.read(1024 * 1024), b''):
            digest.update(chunk)
    return digest.hexdigest()


def _repository_status(repository: Path) -> list[str]:
    result = subprocess.run(
        ['git', 'status', '--short'], cwd=repository, check=True,
        capture_output=True, text=True,
    )
    return result.stdout.splitlines()


def _relative_files(root: Path, directory: str) -> list[str]:
    return [path.relative_to(root).as_posix()
            for path in sorted((root / directory).glob('*')) if path.is_file()]


def _rf_has_internal_gaps(path: Path) -> bool:
    with path.open(newline='') as handle:
        rows = list(csv.reader(handle))[1:]
    serial = []
    for row in rows:
        year, month = map(int, row[0].split('-'))
        serial.append(year * 12 + month)
    return any(right - left != 1 for left, right in zip(serial, serial[1:]))


def _covariance_only_mean_difference(run_root: Path) -> float:
    largest = 0.0
    for specification in 'ABC':
        path = run_root / 'tables' / f'comparison_covariance_only_{specification}_means.csv'
        with path.open(newline='') as handle:
            for row in csv.DictReader(handle):
                largest = max(largest, abs(float(row['markowitz']) - float(row['sim'])))
    return largest


def _field(paragraph, instruction: str, placeholder: str = '') -> None:
    run = paragraph.add_run()
    begin = OxmlElement('w:fldChar')
    begin.set(qn('w:fldCharType'), 'begin')
    separate = OxmlElement('w:fldChar')
    separate.set(qn('w:fldCharType'), 'separate')
    instr = OxmlElement('w:instrText')
    instr.set(qn('xml:space'), 'preserve')
    instr.text = instruction
    end = OxmlElement('w:fldChar')
    end.set(qn('w:fldCharType'), 'end')
    run._r.extend([begin, instr, separate])
    if placeholder:
        value = OxmlElement('w:t')
        value.text = placeholder
        run._r.append(value)
    run._r.append(end)


def _set_repeat_header(row) -> None:
    tr_pr = row._tr.get_or_add_trPr()
    repeat = OxmlElement('w:tblHeader')
    repeat.set(qn('w:val'), 'true')
    tr_pr.append(repeat)


def _prevent_row_split(row) -> None:
    tr_pr = row._tr.get_or_add_trPr()
    tr_pr.append(OxmlElement('w:cantSplit'))


def _shade(cell, fill: str) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    shading = OxmlElement('w:shd')
    shading.set(qn('w:fill'), fill)
    tc_pr.append(shading)


def _set_cell_margins(cell, value: int = 70) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    margins = tc_pr.first_child_found_in('w:tcMar')
    if margins is None:
        margins = OxmlElement('w:tcMar')
        tc_pr.append(margins)
    for edge in ('top', 'left', 'bottom', 'right'):
        node = OxmlElement(f'w:{edge}')
        node.set(qn('w:w'), str(value))
        node.set(qn('w:type'), 'dxa')
        margins.append(node)


def _page_number_footer(section) -> None:
    section.footer.is_linked_to_previous = False
    paragraph = section.footer.paragraphs[0]
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = paragraph.add_run('BUFN741 Assignment 2  •  ')
    run.font.name = 'Arial'
    run.font.size = Pt(8)
    _field(paragraph, 'PAGE', '1')


def _page_number_start(section, value: int = 1) -> None:
    sect_pr = section._sectPr
    node = sect_pr.find(qn('w:pgNumType'))
    if node is None:
        node = OxmlElement('w:pgNumType')
        sect_pr.append(node)
    node.set(qn('w:start'), str(value))


def _set_page(section, landscape: bool = False) -> None:
    # python-docx clones the preceding sectPr when it creates a section.  That
    # includes the TOC section's explicit page-number restart, so remove the
    # cloned restart and let Word continue numbering across orientation changes.
    page_numbering = section._sectPr.find(qn('w:pgNumType'))
    if page_numbering is not None:
        section._sectPr.remove(page_numbering)
    section.orientation = WD_ORIENT.LANDSCAPE if landscape else WD_ORIENT.PORTRAIT
    section.page_width = Inches(11 if landscape else 8.5)
    section.page_height = Inches(8.5 if landscape else 11)
    section.top_margin = Inches(.8)
    section.bottom_margin = Inches(.8)
    section.left_margin = Inches(.75 if landscape else .9)
    section.right_margin = Inches(.75 if landscape else .9)
    section.header_distance = Inches(.3)
    section.footer_distance = Inches(.35)
    _page_number_footer(section)


def _add_inline(paragraph, text: str) -> None:
    parts = re.split(r'(\*\*.*?\*\*|`.*?`)', text)
    for part in parts:
        if not part:
            continue
        if part.startswith('**') and part.endswith('**'):
            run = paragraph.add_run(part[2:-2])
            run.bold = True
        elif part.startswith('`') and part.endswith('`'):
            run = paragraph.add_run(part[1:-1])
            run.font.name = 'Cambria Math'
            run.italic = True
        else:
            paragraph.add_run(part)


def _style_document(document: Document) -> None:
    styles = document.styles
    normal = styles['Normal']
    normal.font.name = 'Arial'
    normal.font.size = Pt(10.5)
    normal.paragraph_format.space_after = Pt(5)
    normal.paragraph_format.line_spacing = 1.12
    for name, size, color in [('Title', 24, '17365D'), ('Subtitle', 15, '44546A'),
                              ('Heading 1', 16, '17365D'), ('Heading 2', 12.5, '2F5597')]:
        style = styles[name]
        style.font.name = 'Arial'
        style.font.size = Pt(size)
        style.font.color.rgb = RGBColor.from_string(color)
        style.font.bold = name != 'Subtitle'
        style.paragraph_format.keep_with_next = True
        style.paragraph_format.space_before = Pt(10)
        style.paragraph_format.space_after = Pt(5)
    styles['Heading 1'].paragraph_format.page_break_before = True
    caption = styles['Caption']
    caption.font.name = 'Arial'
    caption.font.size = Pt(9)
    caption.font.bold = True
    caption.font.color.rgb = RGBColor.from_string('404040')
    caption.paragraph_format.keep_with_next = True
    if 'Equation' not in [style.name for style in styles]:
        equation = styles.add_style('Equation', WD_STYLE_TYPE.PARAGRAPH)
        equation.font.name = 'Cambria Math'
        equation.font.size = Pt(11)
        equation.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.CENTER
        equation.paragraph_format.space_before = Pt(5)
        equation.paragraph_format.space_after = Pt(7)


def _add_title_and_toc(document: Document, report_date: str) -> None:
    section = document.sections[0]
    _set_page(section)
    section.footer.paragraphs[0].clear()
    paragraph = document.add_paragraph()
    paragraph.paragraph_format.space_before = Pt(120)
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = paragraph.add_run('BUFN741 GROUP ASSIGNMENT 2')
    run.bold = True
    run.font.name = 'Arial'
    run.font.size = Pt(26)
    run.font.color.rgb = RGBColor.from_string('17365D')
    subtitle = document.add_paragraph()
    subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
    subtitle.paragraph_format.space_before = Pt(18)
    run = subtitle.add_run('Multi-Asset Portfolio Optimization:\nMarkowitz and Single-Index Models')
    run.font.name = 'Arial'
    run.font.size = Pt(17)
    run.font.color.rgb = RGBColor.from_string('2F5597')
    course = document.add_paragraph()
    course.alignment = WD_ALIGN_PARAGRAPH.CENTER
    course.paragraph_format.space_before = Pt(36)
    course.add_run('BUFN741 Advanced Capital Markets\nGroup Assignment\n\n' + report_date)
    toc_section = document.add_section(WD_SECTION.NEW_PAGE)
    _set_page(toc_section)
    _page_number_start(toc_section, 1)
    title = document.add_paragraph('Table of Contents', style='Title')
    title.alignment = WD_ALIGN_PARAGRAPH.LEFT
    toc = document.add_paragraph()
    _field(toc, r'TOC \o "1-2" \h \z \u', 'Update fields to build the table of contents.')
    document.add_page_break()


def _parse_table(lines: list[str]) -> list[list[str]]:
    rows = []
    for line in lines:
        rows.append([cell.strip() for cell in line.strip().strip('|').split('|')])
    return [rows[0]] + rows[2:]


def _add_table(document: Document, rows: list[list[str]], compact: bool) -> None:
    table = document.add_table(rows=len(rows), cols=len(rows[0]))
    table.style = 'Table Grid'
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = True
    for row_index, values in enumerate(rows):
        row = table.rows[row_index]
        _prevent_row_split(row)
        if row_index == 0:
            _set_repeat_header(row)
        for column_index, value in enumerate(values):
            cell = row.cells[column_index]
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            _set_cell_margins(cell, 45 if compact else 70)
            paragraph = cell.paragraphs[0]
            paragraph.paragraph_format.space_after = Pt(0)
            paragraph.alignment = (WD_ALIGN_PARAGRAPH.LEFT if column_index < 2
                                   else WD_ALIGN_PARAGRAPH.RIGHT)
            run = paragraph.add_run(value)
            run.font.name = 'Arial'
            run.font.size = Pt(7.5 if compact else 8.5)
            if row_index == 0:
                run.bold = True
                run.font.color.rgb = RGBColor(255, 255, 255)
                _shade(cell, '2F5597')
            elif row_index % 2 == 0:
                _shade(cell, 'EAF0F8')
    document.add_paragraph().paragraph_format.space_after = Pt(1)


def _add_equations_after_heading(document: Document, heading: str) -> None:
    if heading == '1. Objective and Methodology':
        for equation in ['E(rₚ) = w′μ', 'Var(rₚ) = w′Σw', 'Sharpe(rₚ) = [E(rₚ) − r_f] / σₚ']:
            document.add_paragraph(equation, style='Equation')
    elif heading == '8. Single-Index Model Estimation':
        document.add_paragraph('rᵢ − r_f = αᵢ + βᵢ(r_SPY − r_f) + εᵢ', style='Equation')
        document.add_paragraph('Σ_SIM = ββ′ Var(r_SPY − r_f) + diag(σ²_ε)', style='Equation')


def build_docx(source: Path, output: Path, run_root: Path, pytest_result: str,
               report_date: str) -> None:
    document = Document()
    _style_document(document)
    _add_title_and_toc(document, report_date)
    lines = source.read_text().splitlines()
    start = next(i for i, line in enumerate(lines) if line == '## Executive Summary')
    index, landscape = start, False
    pending_table_number = None
    while index < len(lines):
        line = lines[index].rstrip()
        if not line:
            index += 1
            continue
        if line.startswith('## '):
            heading = line[3:]
            document.add_paragraph(heading, style='Heading 1')
            _add_equations_after_heading(document, heading)
        elif line.startswith('### '):
            document.add_paragraph(line[4:], style='Heading 2')
        elif re.fullmatch(r'\*\*Table \d+\..+\*\*', line):
            pending_table_number = int(re.search(r'Table (\d+)', line).group(1))
            if pending_table_number in LANDSCAPE_TABLES and not landscape:
                section = document.add_section(WD_SECTION.NEW_PAGE)
                _set_page(section, True)
                landscape = True
            paragraph = document.add_paragraph(style='Caption')
            paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
            paragraph.add_run(line[2:-2])
        elif line.startswith('|'):
            table_lines = []
            while index < len(lines) and lines[index].startswith('|'):
                table_lines.append(lines[index])
                index += 1
            rows = _parse_table(table_lines)
            _add_table(document, rows, len(rows[0]) >= 8)
            if landscape:
                section = document.add_section(WD_SECTION.NEW_PAGE)
                _set_page(section, False)
                landscape = False
            pending_table_number = None
            continue
        elif line.startswith('!['):
            match = re.fullmatch(r'!\[(.+)\]\((.+)\)', line)
            if not match:
                raise ValueError(f'Invalid image syntax: {line}')
            caption, relative = match.groups()
            image = run_root / relative
            if not image.is_file():
                raise FileNotFoundError(image)
            paragraph = document.add_paragraph()
            paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
            paragraph.paragraph_format.keep_with_next = True
            paragraph.add_run().add_picture(str(image), width=Inches(6.25))
            cap = document.add_paragraph(caption, style='Caption')
            cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
            # The Markdown has the same caption on the next non-empty line.
            lookahead = index + 1
            while lookahead < len(lines) and not lines[lookahead].strip():
                lookahead += 1
            if lookahead < len(lines) and lines[lookahead].strip() == f'*{caption}*':
                index = lookahead
        elif line.startswith('- '):
            content = line[2:]
            if content.startswith('Pytest result:'):
                content = f'Pytest result: {pytest_result}.'
            paragraph = document.add_paragraph(style='List Bullet')
            _add_inline(paragraph, content)
        elif line.startswith('*') and line.endswith('*'):
            # Standalone figure captions are consumed with the image.
            pass
        else:
            paragraph = document.add_paragraph()
            _add_inline(paragraph, line)
        index += 1

    settings = document.settings._element
    update = settings.find(qn('w:updateFields'))
    if update is None:
        update = OxmlElement('w:updateFields')
        settings.append(update)
    update.set(qn('w:val'), 'true')
    document.core_properties.title = 'BUFN741 Group Assignment 2'
    document.core_properties.subject = 'Multi-Asset Portfolio Optimization: Markowitz and Single-Index Models'
    document.core_properties.comments = 'Generated from validated frozen production-run Markdown.'
    document.save(output)


def _uno_export(docx: Path, pdf: Path, profile: Path) -> None:
    port = 22000 + (os.getpid() % 10000)
    command = ['libreoffice', f'-env:UserInstallation=file://{profile}', '--headless', '--nologo',
               '--nodefault', '--nofirststartwizard',
               f'--accept=socket,host=127.0.0.1,port={port};urp;StarOffice.ComponentContext']
    process = subprocess.Popen(command, stdout=subprocess.PIPE, stderr=subprocess.PIPE)
    try:
        sys.path.append('/usr/lib/python3/dist-packages')
        import uno
        from com.sun.star.beans import PropertyValue

        def prop(name: str, value):
            item = PropertyValue()
            item.Name, item.Value = name, value
            return item

        context = None
        for _ in range(100):
            try:
                local = uno.getComponentContext()
                resolver = local.ServiceManager.createInstanceWithContext(
                    'com.sun.star.bridge.UnoUrlResolver', local)
                context = resolver.resolve(
                    f'uno:socket,host=127.0.0.1,port={port};urp;StarOffice.ComponentContext')
                break
            except Exception:
                time.sleep(.1)
        if context is None:
            raise RuntimeError('Could not connect to headless LibreOffice')
        desktop = context.ServiceManager.createInstanceWithContext('com.sun.star.frame.Desktop', context)
        document = desktop.loadComponentFromURL(docx.resolve().as_uri(), '_blank', 0, (prop('Hidden', True),))
        if document is None:
            raise RuntimeError('LibreOffice could not open generated DOCX')
        indexes = document.getDocumentIndexes()
        for idx in range(indexes.getCount()):
            indexes.getByIndex(idx).update()
        document.store()
        document.storeToURL(pdf.resolve().as_uri(), (prop('FilterName', 'writer_pdf_Export'),))
        document.close(True)
        desktop.terminate()
    finally:
        process.terminate()
        try:
            process.wait(timeout=10)
        except subprocess.TimeoutExpired:
            process.kill()


def _pdf_pages(pdf: Path) -> int:
    output = subprocess.run(['pdfinfo', str(pdf)], check=True, capture_output=True, text=True).stdout
    match = re.search(r'^Pages:\s+(\d+)$', output, re.MULTILINE)
    if not match:
        raise RuntimeError('pdfinfo did not report a page count')
    return int(match.group(1))


def validate_production(build: Path, run_root: Path, source_text: str) -> dict:
    docx = build / 'final_assignment_report.docx'
    pdf = build / 'final_assignment_report.pdf'
    markdown = build / 'final_assignment_report.md'
    for path in [docx, pdf, markdown]:
        if not path.is_file() or path.stat().st_size == 0:
            raise RuntimeError(f'Missing or empty production artifact: {path.name}')
    with zipfile.ZipFile(docx) as archive:
        bad = archive.testzip()
        if bad:
            raise RuntimeError(f'Corrupt DOCX entry: {bad}')
        document_xml = archive.read('word/document.xml').decode('utf-8')
        footer_xml = ''.join(archive.read(name).decode('utf-8') for name in archive.namelist()
                             if name.startswith('word/footer') and name.endswith('.xml'))
        media = [name for name in archive.namelist() if name.startswith('word/media/')]
    reopened = Document(docx)
    if len(reopened.inline_shapes) != 15 or len(media) != 15:
        raise RuntimeError('DOCX does not contain exactly 15 embedded figures')
    if len(reopened.tables) != 19:
        raise RuntimeError(f'DOCX contains {len(reopened.tables)} tables; expected 19')
    if 'TOC' not in document_xml or ('PAGE' not in footer_xml and 'pgNum' not in footer_xml):
        raise RuntimeError('DOCX is missing TOC or page-number fields')
    visible = '\n'.join(paragraph.text for paragraph in reopened.paragraphs)
    for section in REQUIRED_SECTIONS:
        if section not in visible:
            raise RuntimeError(f'DOCX missing section: {section}')
    for value in ['2006-05 through 2026-08, N=244', '2006-05 through 2026-07, N=243', 'NONATTAINED']:
        if value not in visible:
            raise RuntimeError(f'DOCX missing authoritative content: {value}')
    page_count = _pdf_pages(pdf)
    if not 15 <= page_count <= 100:
        raise RuntimeError(f'Implausible PDF page count: {page_count}')
    text_path = build / '.pdf_text.txt'
    subprocess.run(['pdftotext', str(pdf), str(text_path)], check=True)
    pdf_text = text_path.read_text(errors='replace')
    text_path.unlink()
    for section in REQUIRED_SECTIONS:
        if section not in pdf_text:
            raise RuntimeError(f'PDF missing section: {section}')
    if '![' in pdf_text or re.search(r'^#{1,4}\s', pdf_text, re.MULTILINE):
        raise RuntimeError('PDF contains unresolved Markdown syntax')
    render_dir = build / '.render_check'
    render_dir.mkdir()
    subprocess.run(['pdftoppm', '-r', '20', '-png', str(pdf), str(render_dir / 'page')],
                   check=True, stdout=subprocess.DEVNULL, stderr=subprocess.PIPE)
    pages = sorted(render_dir.glob('page-*.png'))
    if len(pages) != page_count or any(path.stat().st_size == 0 for path in pages):
        raise RuntimeError('Not every PDF page rendered successfully')
    shutil.rmtree(render_dir)
    image_refs = re.findall(r'!\[[^]]+\]\(([^)]+)\)', markdown.read_text())
    if len(image_refs) != 15 or any(not (build / ref).resolve().is_file() for ref in image_refs):
        raise RuntimeError('Production Markdown figure references are incomplete')
    if source_text.count('NONATTAINED') != markdown.read_text().count('NONATTAINED'):
        raise RuntimeError('Production Markdown changed analytical status content')
    return {'pdf_pages': page_count, 'docx_tables': len(reopened.tables),
            'docx_figures': len(reopened.inline_shapes), 'required_sections': len(REQUIRED_SECTIONS)}


def generate(run_root: Path, pytest_result: str) -> Path:
    run_root = run_root.resolve()
    repository = run_root.parent.parent
    source = run_root / 'final_assignment_report.md'
    if not source.is_file():
        raise FileNotFoundError(source)
    state = json.loads((run_root / 'state.json').read_text())
    if state['status'] != 'PASSED' or not all(stage['status'] == 'PASSED' for stage in state['stages'].values()):
        raise RuntimeError('Production requires a fully PASSED analytical run')
    if state.get('attempt') != 1 or state.get('attempt_history'):
        raise RuntimeError('Canonical production requires a fresh, non-resumed analytical run')
    production = run_root / 'production'
    if production.exists() or (production / 'FINALIZED').exists():
        raise FileExistsError('Production is already finalized; use a new production run/version')
    build = run_root / f'.production_build_{os.getpid()}'
    if build.exists():
        raise FileExistsError(build)
    build.mkdir()
    try:
        now = datetime.now(timezone.utc)
        # Selecting a validated run for canonical packaging is production
        # metadata only; it does not alter any analytical artifact.
        state['run_type'] = 'production'
        state['production_requested_at'] = now.isoformat()
        state_path = run_root / 'state.json'
        state_temp = run_root / '.state.production.tmp'
        state_temp.write_text(json.dumps(state, indent=2) + '\n')
        state_temp.replace(state_path)
        source_text = source.read_text()
        packaged_text = source_text.replace('](figures/', '](../figures/')
        markdown = build / 'final_assignment_report.md'
        markdown.write_text(packaged_text)
        docx = build / 'final_assignment_report.docx'
        pdf = build / 'final_assignment_report.pdf'
        build_docx(source, docx, run_root, pytest_result, now.date().isoformat())
        profile = build / '.libreoffice_profile'
        profile.mkdir()
        _uno_export(docx, pdf, profile)
        shutil.rmtree(profile, ignore_errors=True)
        validation = validate_production(build, run_root, source_text)
        environment = json.loads((run_root / 'diagnostics' / 'environment.json').read_text())
        summary = json.loads((run_root / 'results_summary.json').read_text())
        artifact_manifest_path = run_root / 'diagnostics' / 'artifact_manifest.json'
        artifact_manifest = json.loads(artifact_manifest_path.read_text())
        repository_status = _repository_status(repository)
        config_path = repository / 'config' / 'assignment2.yaml'
        snapshot_path = run_root / 'config_snapshot.yaml'
        rf_path = run_root / 'data' / 'fama_french_rf_available.csv'
        reports = [name for name in [
            'markowitz_report.md', 'sim_report.md', 'model_comparison_report.md',
            'final_assignment_report.md',
        ] if (run_root / name).is_file()]
        manifest = {
            'production_run_id': state['run_id'],
            'run_type': 'production',
            'generation_timestamp_utc': now.isoformat(),
            'git_commit': environment.get('git_commit'),
            'repository_status': {
                'clean': not repository_status,
                'porcelain': repository_status,
            },
            'config': {
                'path': 'config/assignment2.yaml',
                'checksum_sha256': sha256(config_path),
                'snapshot_path': '../config_snapshot.yaml',
                'snapshot_checksum_sha256': sha256(snapshot_path),
            },
            'config_checksum_sha256': sha256(snapshot_path),
            'pytest_result': pytest_result,
            'samples': summary['samples'],
            'rf_lag': {
                'trailing_months': summary['sample_policy']['trailing_rf_lag_months'],
                'accepted': summary['sample_policy']['lag_accepted'],
                'allowed': summary['sample_policy']['allow_trailing_rf_lag'],
                'maximum_months': summary['sample_policy']['max_trailing_rf_lag_months'],
                'internal_gaps': _rf_has_internal_gaps(rf_path),
            },
            'workflow': {
                'status': state['status'],
                'stages': {name: details['status'] for name, details in state['stages'].items()},
            },
            'covariance_only_max_abs_mean_difference': _covariance_only_mean_difference(run_root),
            'source_markdown_path': '../final_assignment_report.md',
            'docx_path': 'final_assignment_report.docx',
            'pdf_path': 'final_assignment_report.pdf',
            'principal_outputs': {
                'reports': reports + [
                    'production/final_assignment_report.md',
                    'production/final_assignment_report.docx',
                    'production/final_assignment_report.pdf',
                ],
                'tables': _relative_files(run_root, 'tables'),
                'figures': _relative_files(run_root, 'figures'),
            },
            'validation': validation,
            'files': {path.name: {'sha256': sha256(path), 'bytes': path.stat().st_size}
                      for path in [markdown, docx, pdf]},
            'authoritative_results': {
                'path': '../results_summary.json',
                'sha256': sha256(run_root / 'results_summary.json'),
                'bytes': (run_root / 'results_summary.json').stat().st_size,
            },
            'analytical_artifact_manifest': {
                'path': '../diagnostics/artifact_manifest.json',
                'sha256': sha256(artifact_manifest_path),
            },
            'data_cache_checksums_sha256': {
                path: checksum for path, checksum in artifact_manifest.items()
                if path.startswith('data/raw/')
            },
        }
        manifest_path = build / 'production_manifest.json'
        manifest_path.write_text(json.dumps(manifest, indent=2) + '\n')
        marker = build / 'FINALIZED'
        marker.write_text(f"Finalized {now.isoformat()}\nManifest SHA-256: {sha256(manifest_path)}\n")
        build.rename(production)
        return production
    except BaseException:
        shutil.rmtree(build, ignore_errors=True)
        raise


def main() -> int:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument('--run', required=True, type=Path, help='Validated frozen run directory')
    parser.add_argument('--pytest-result', required=True, help='Exact completed-suite result for provenance')
    args = parser.parse_args()
    output = generate(args.run, args.pytest_result)
    print(f'Production finalized: {output}')
    return 0


if __name__ == '__main__':
    raise SystemExit(main())
